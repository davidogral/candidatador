from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

from pypdf import PdfReader
from sqlmodel import Session, select

from candidatador.config import Paths
from candidatador.models import Document, DocumentKind
from candidatador.sources.base import matches_keywords

TITLE_WEIGHT = 3
TEXT_SUFFIXES = {".txt", ".md", ".markdown"}


def extract_text(path: Path) -> str:
    """Best-effort plain-text extraction. Images and unknown formats return ''."""
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
        if suffix == ".docx":
            from docx import Document as DocxDocument

            return "\n".join(p.text for p in DocxDocument(str(path)).paragraphs).strip()
        if suffix in TEXT_SUFFIXES:
            return path.read_text(encoding="utf-8", errors="ignore").strip()
    except Exception:  # corrupted/encrypted files shouldn't break the upload
        return ""
    return ""


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def add_document(
    s: Session,
    paths: Paths,
    source: Path,
    *,
    kind: DocumentKind,
    title: str | None = None,
    language: str = "pt",
    tags: list[str] | None = None,
    is_default: bool = False,
    issued_by: str = "",
    issued_at: str = "",
) -> Document:
    """Copy a file into the vault and register it. Re-adding the same file is a no-op."""
    source = source.expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)

    digest = _sha256(source)
    existing = s.exec(select(Document).where(Document.sha256 == digest)).first()
    if existing:
        return existing

    paths.ensure()
    dest_dir = paths.vault / kind.value
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / f"{digest[:12]}{source.suffix.lower()}"
    shutil.copy2(source, dest)

    if is_default:
        for doc in s.exec(select(Document).where(Document.kind == kind, Document.is_default)):
            doc.is_default = False
            s.add(doc)

    doc = Document(
        kind=kind,
        title=title or source.stem,
        language=language,
        tags_json=json.dumps(tags or []),
        path=str(dest),
        original_name=source.name,
        sha256=digest,
        text=extract_text(dest),
        is_default=is_default,
        issued_by=issued_by,
        issued_at=issued_at,
    )
    s.add(doc)
    s.commit()
    s.refresh(doc)
    return doc


def pick_resume(s: Session, job_text: str, *, language: str | None = None) -> Document | None:
    """Choose the resume version whose tags best overlap the job text.

    Falls back to the default resume, then to the most recent one. The AI matcher can
    override this choice when enabled.
    """
    stmt = select(Document).where(Document.kind == DocumentKind.RESUME)
    resumes = list(s.exec(stmt))
    if language:
        resumes = [r for r in resumes if r.language == language] or resumes
    if not resumes:
        return None

    # The first line is the job title: a tag there says more about the role than the
    # description does ("Analista de BI" whose description also lists ML tools).
    title = job_text.split("\n", 1)[0]

    def rank(doc: Document) -> tuple[int, bool, float]:
        # whole words only: a "bi" tag must not match "ambiente" or "habilidades"
        hits = sum(
            (TITLE_WEIGHT if matches_keywords(title, [tag]) else 0)
            + (1 if matches_keywords(job_text, [tag]) else 0)
            for tag in doc.tags
        )
        return hits, doc.is_default, doc.created_at.timestamp()

    return max(resumes, key=rank)
